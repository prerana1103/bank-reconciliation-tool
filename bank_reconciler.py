#!/usr/bin/env python3
"""
bank_reconciler.py - Automated Bank Reconciliation Engine
==========================================================

Matches bank-statement transactions against accounting (cash-book / GL) records,
auto-classifies every discrepancy, and exports an audit-ready Excel workbook with
a sign-off exception report and a tamper-evident run log.

Matching waterfall (each rule only sees what earlier rules left unmatched):
    R1  Reference + Amount (+ date window)             -> clean match
    R2  Amount + Date window + Description similarity   -> clean match
    R3  Same reference, DIFFERENT amount                -> AMOUNT_MISMATCH
    R4  Amount + Date window only (no corroboration)    -> LOW_CONFIDENCE_MATCH
    R5  One-to-many / many-to-one (batch deposits)      -> SPLIT_MATCH
    R6  Similar description + date, near amount         -> AMOUNT_MISMATCH
Residual items are then classified as DUPLICATE_* or UNMATCHED_*.

All money is held as integer cents internally - no floating-point drift.
"""
from __future__ import annotations

import argparse
import bisect
import getpass
import hashlib
import json
import platform
import re
import sys
import uuid
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from difflib import SequenceMatcher
from functools import lru_cache
from itertools import combinations
from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

__version__ = "1.0.0"

# --------------------------------------------------------------------------- #
# Constants
# --------------------------------------------------------------------------- #
UNMATCHED_BANK = "UNMATCHED_BANK"
UNMATCHED_LEDGER = "UNMATCHED_LEDGER"
DUPLICATE_BANK = "DUPLICATE_BANK"
DUPLICATE_LEDGER = "DUPLICATE_LEDGER"
AMOUNT_MISMATCH = "AMOUNT_MISMATCH"
LOW_CONFIDENCE = "LOW_CONFIDENCE_MATCH"
SPLIT_MATCH = "SPLIT_MATCH"
TYPE_ORDER = [UNMATCHED_BANK, UNMATCHED_LEDGER, DUPLICATE_BANK, DUPLICATE_LEDGER,
              AMOUNT_MISMATCH, LOW_CONFIDENCE, SPLIT_MATCH]

R1 = "R1 Reference + Amount"
R2 = "R2 Amount + Date + Description"
R3 = "R3 Reference match, amount differs"
R4 = "R4 Amount + Date only"
R5 = "R5 Split / batch (1:N)"
R6 = "R6 Fuzzy description, amount differs"

LEVELS = ["Low", "Medium", "High"]
REF_PREFIXES = "INV|CHQ|CHK|PO|REF|UTR|TXN|BILL"
STOPWORDS = {"PAYMENT", "PMT", "RECEIPT", "TRANSFER", "NEFT", "RTGS", "IMPS", "UPI",
             "ACH", "WIRE", "TO", "FROM", "BY", "THE", "OF", "FOR", "AND", "REF",
             "INVOICE", "CHEQUE", "DEPOSIT", "BILL"}
DISPOSITIONS = ("Accepted - no action,Adjusting entry required,Investigate with bank,"
                "Reverse duplicate entry,Write-off approved")

ALIASES = {
    "date": ["date", "transaction date", "txn date", "posting date", "value date",
             "booking date", "trans date"],
    "description": ["description", "narration", "details", "memo", "particulars",
                    "transaction details", "payee", "remarks"],
    "amount": ["amount", "amt", "net amount", "transaction amount"],
    "debit": ["debit", "withdrawal", "withdrawals", "dr", "paid out", "money out"],
    "credit": ["credit", "deposit", "deposits", "cr", "paid in", "money in"],
    "reference": ["reference", "ref", "ref no", "reference number", "cheque no",
                  "check no", "chq no", "utr", "invoice", "document no", "doc no",
                  "transaction id", "txn id"],
}


# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #
@dataclass
class Config:
    date_tolerance_days: int = 3          # bank vs ledger posting lag for non-reference rules
    ref_date_window_days: int = 15        # wider window when a reference corroborates (cheques)
    amount_tolerance: float = 0.0         # absolute tolerance treated as "equal"
    desc_match_threshold: float = 0.70    # R2 similarity floor
    mismatch_desc_threshold: float = 0.75  # R6 similarity floor
    mismatch_max_pct: float = 0.25        # R6 max |diff| as share of the larger amount
    rounding_threshold: float = 1.00      # diffs up to this are classed ROUNDING_DIFFERENCE
    split_max_items: int = 3
    split_date_window_days: int = 5
    split_candidate_cap: int = 25
    duplicate_window_days: int = 1
    outstanding_days: int = 5             # unmatched ledger items this close to period end = in transit
    high_materiality: float = 10000.0
    medium_materiality: float = 1000.0
    dayfirst: bool = False
    bank_sign: int = 1                    # set -1 to invert a statement's sign convention
    ledger_sign: int = 1


# --------------------------------------------------------------------------- #
# Data model
# --------------------------------------------------------------------------- #
@dataclass(eq=False)
class Txn:
    uid: str
    source: str            # "BANK" | "LEDGER"
    row: int               # row number in the source file
    date: date
    desc: str
    ref: str
    amt: int               # cents; + = money in, - = money out
    ndesc: str = ""
    nref: str = ""


@dataclass(eq=False)
class Match:
    match_id: str
    rule: str
    score: float
    bank: list
    ledger: list
    kind: str = "MATCH"    # MATCH | AMOUNT_MISMATCH | LOW_CONFIDENCE | SPLIT
    subtype: str = ""

    @property
    def bank_total(self) -> int:
        return sum(t.amt for t in self.bank)

    @property
    def ledger_total(self) -> int:
        return sum(t.amt for t in self.ledger)


@dataclass
class ExceptionItem:
    exc_id: str = ""
    type: str = ""
    subtype: str = ""
    severity: str = "Low"
    bank_ids: str = ""
    ledger_ids: str = ""
    bank_date: Optional[date] = None
    ledger_date: Optional[date] = None
    reference: str = ""
    description: str = ""
    bank_amt: Optional[int] = None
    ledger_amt: Optional[int] = None
    explanation: str = ""
    action: str = ""
    match_id: str = ""

    @property
    def impact(self) -> int:
        return (self.bank_amt or 0) - (self.ledger_amt or 0)


# --------------------------------------------------------------------------- #
# Normalisation helpers
# --------------------------------------------------------------------------- #
def parse_amount(value) -> Optional[int]:
    """Parse messy amount input into integer cents. Returns None if unparseable."""
    if value is None:
        return None
    if isinstance(value, float) and pd.isna(value):
        return None
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        d = Decimal(str(value))
    else:
        s = str(value).strip()
        if not s or s.lower() in ("nan", "none", "-", "<na>"):
            return None
        neg, u = False, s.upper()
        if u.endswith("DR"):
            neg, s = True, s[:-2]
        elif u.endswith("CR"):
            s = s[:-2]
        s = s.strip()
        if s.startswith("(") and s.endswith(")"):
            neg, s = True, s[1:-1]
        if s.endswith("-"):
            neg, s = True, s[:-1]
        s = re.sub(r"[^\d.\-]", "", s)
        if s in ("", "-", "."):
            return None
        try:
            d = Decimal(s)
        except InvalidOperation:
            return None
        if neg:
            d = -abs(d)
    return int((d * 100).to_integral_value(rounding=ROUND_HALF_UP))


def norm_text(s: str) -> str:
    s = re.sub(r"[^A-Z0-9]+", " ", (s or "").upper())
    s = re.sub(rf"\b({REF_PREFIXES})\s*(\d)", r"\1\2", s)   # "INV 1001" -> "INV1001"
    return re.sub(r"\s+", " ", s).strip()


def norm_ref(s: str) -> str:
    s = re.sub(r"[^A-Z0-9]", "", (s or "").upper())
    return re.sub(rf"^({REF_PREFIXES})", "", s)


def extract_ref(ndesc: str) -> str:
    m = re.search(rf"\b(?:{REF_PREFIXES})(\d{{3,}})\b", ndesc)
    return m.group(1) if m else ""


def is_doc_ref(nref: str) -> bool:
    return bool(re.fullmatch(r"\d{3,}", nref or ""))


@lru_cache(maxsize=500_000)
def similarity(a: str, b: str) -> float:
    """Blend of character-level and token-level similarity, 0..1."""
    if not a or not b:
        return 0.0
    seq = SequenceMatcher(None, a, b).ratio()
    ta = {t for t in a.split() if t not in STOPWORDS and len(t) > 1}
    tb = {t for t in b.split() if t not in STOPWORDS and len(t) > 1}
    tok = 0.0
    if ta and tb:
        inter = len(ta & tb)
        tok = 0.5 * inter / len(ta | tb) + 0.5 * inter / min(len(ta), len(tb))
    return max(seq, tok)


def fm(cents: int) -> str:
    return f"{cents / 100:,.2f}"


# --------------------------------------------------------------------------- #
# Loading
# --------------------------------------------------------------------------- #
def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _resolve_columns(df: pd.DataFrame, overrides: dict, label: str) -> dict:
    lookup = {str(c).strip().lower(): c for c in df.columns}
    cols = {}
    for fld, aliases in ALIASES.items():
        if fld in overrides:
            key = overrides[fld].strip().lower()
            if key not in lookup:
                raise ValueError(f"[{label}] column '{overrides[fld]}' not found. "
                                 f"Available: {list(df.columns)}")
            cols[fld] = lookup[key]
            continue
        for a in aliases:
            if a in lookup:
                cols[fld] = lookup[a]
                break
    if "date" not in cols:
        raise ValueError(f"[{label}] no date column found. Available: {list(df.columns)}. "
                         f"Use --{label.lower()}-cols date=<column>")
    if "amount" not in cols and not ("debit" in cols or "credit" in cols):
        raise ValueError(f"[{label}] no amount (or debit/credit) columns found. "
                         f"Available: {list(df.columns)}")
    return cols


def load_transactions(path, source: str, cfg: Config, overrides: Optional[dict] = None,
                      sheet=0):
    """Return (transactions, rejected_rows, file_info). Never silently drops data."""
    path = Path(path)
    overrides = overrides or {}
    if path.suffix.lower() in (".xlsx", ".xlsm", ".xls"):
        df = pd.read_excel(path, sheet_name=sheet, dtype=object)
    else:
        df = pd.read_csv(path, dtype=str)
    cols = _resolve_columns(df, overrides, source)

    try:
        dates = pd.to_datetime(df[cols["date"]], dayfirst=cfg.dayfirst,
                               errors="coerce", format="mixed")
    except (TypeError, ValueError):
        dates = pd.to_datetime(df[cols["date"]], dayfirst=cfg.dayfirst, errors="coerce")

    # Convention: bank debit = money OUT; ledger (cash book) debit = money IN.
    debit_is_inflow = source == "LEDGER"
    sign = cfg.bank_sign if source == "BANK" else cfg.ledger_sign
    prefix = "B" if source == "BANK" else "L"

    def cell(r, key):
        if key not in cols:
            return ""
        v = r[cols[key]]
        return "" if (v is None or (isinstance(v, float) and pd.isna(v))) else str(v).strip()

    txns, rejects, n = [], [], 0
    for i, (_, r) in enumerate(df.iterrows()):
        rownum = i + 2
        if all(str(v).strip() in ("", "nan", "NaT", "None") for v in r.values):
            continue
        d = dates.iloc[i]
        if pd.isna(d):
            rejects.append({"source": source, "row": rownum, "reason": "Unparseable date",
                            "raw": " | ".join(str(v) for v in r.values)})
            continue
        if "amount" in cols:
            a = parse_amount(r[cols["amount"]])
        else:
            dr = parse_amount(r[cols["debit"]]) if "debit" in cols else None
            cr = parse_amount(r[cols["credit"]]) if "credit" in cols else None
            if dr is None and cr is None:
                a = None
            else:
                inflow, outflow = (dr, cr) if debit_is_inflow else (cr, dr)
                a = (inflow or 0) - (outflow or 0)
        if a is None:
            rejects.append({"source": source, "row": rownum, "reason": "Unparseable amount",
                            "raw": " | ".join(str(v) for v in r.values)})
            continue
        n += 1
        desc, ref = cell(r, "description"), cell(r, "reference")
        nd = norm_text(desc)
        txns.append(Txn(uid=f"{prefix}{n:05d}", source=source, row=rownum, date=d.date(),
                        desc=desc, ref=ref, amt=a * sign, ndesc=nd,
                        nref=norm_ref(ref) or extract_ref(nd)))
    info = {"source": source, "file": path.name, "path": str(path.resolve()),
            "bytes": path.stat().st_size, "sha256": sha256_file(path),
            "rows_read": len(df), "rows_loaded": len(txns), "rows_rejected": len(rejects),
            "columns_used": {k: str(v) for k, v in cols.items()}}
    return txns, rejects, info


# --------------------------------------------------------------------------- #
# Reconciliation engine
# --------------------------------------------------------------------------- #
class Reconciler:
    def __init__(self, bank: list, ledger: list, cfg: Optional[Config] = None):
        self.bank, self.ledger = bank, ledger
        self.cfg = cfg or Config()
        self.tol_c = int(round(self.cfg.amount_tolerance * 100))
        self.matches: list[Match] = []
        self.exceptions: list[ExceptionItem] = []
        self.rule_log: list[tuple[str, int]] = []
        self.matched: dict[str, Match] = {}
        self.txn_exc: dict[str, str] = {}
        self.period_end: Optional[date] = max((t.date for t in bank), default=None)

    # -- public ------------------------------------------------------------- #
    def run(self) -> "Reconciler":
        self._r1_reference()
        self._r2_amount_date_desc()
        self._r3_reference_mismatch()
        self._r4_amount_date_only()
        self._r5_split()
        self._r6_fuzzy_mismatch()
        self._classify()
        return self

    def status(self, t: Txn) -> str:
        m = self.matched.get(t.uid)
        if m is None:
            return "UNMATCHED"
        return "MATCHED" if m.kind == "MATCH" else "MATCHED-REVIEW"

    def counts_by_type(self) -> dict:
        return dict(Counter(e.type for e in self.exceptions))

    # -- helpers ------------------------------------------------------------ #
    def _free(self, txns):
        return [t for t in txns if t.uid not in self.matched]

    def _amt_eq(self, a: Txn, b: Txn) -> bool:
        return abs(a.amt - b.amt) <= self.tol_c

    @staticmethod
    def _dd(a: Txn, b: Txn) -> int:
        return abs((a.date - b.date).days)

    @staticmethod
    def _ref_conflict(b: Txn, l: Txn) -> bool:
        return (is_doc_ref(b.nref) and is_doc_ref(l.nref) and b.nref != l.nref)

    def _commit(self, bank, ledger, rule, score, kind="MATCH", subtype="") -> Match:
        m = Match(f"M{len(self.matches) + 1:05d}", rule, round(score, 3),
                  list(bank), list(ledger), kind, subtype)
        for t in m.bank + m.ledger:
            self.matched[t.uid] = m
        self.matches.append(m)
        return m

    def _assign(self, cands, rule, kind="MATCH", subtype_fn=None):
        cands.sort(key=lambda c: (-c[0], c[1].uid, c[2].uid))
        n = 0
        for score, b, l in cands:
            if b.uid in self.matched or l.uid in self.matched:
                continue
            self._commit([b], [l], rule, score, kind, subtype_fn(b, l) if subtype_fn else "")
            n += 1
        self.rule_log.append((rule, n))

    # -- rules -------------------------------------------------------------- #
    def _ref_index(self):
        idx = defaultdict(list)
        for l in self._free(self.ledger):
            if l.nref:
                idx[l.nref].append(l)
        return idx

    def _r1_reference(self):
        idx, c = self._ref_index(), []
        for b in self._free(self.bank):
            for l in idx.get(b.nref, []) if b.nref else []:
                dd = self._dd(b, l)
                if dd <= self.cfg.ref_date_window_days and self._amt_eq(b, l):
                    c.append((max(0.9, 1 - 0.005 * dd), b, l))
        self._assign(c, R1)

    def _r2_amount_date_desc(self):
        idx, c = defaultdict(list), []
        for l in self._free(self.ledger):
            idx[l.amt].append(l)
        w = self.cfg.date_tolerance_days
        for b in self._free(self.bank):
            for a in range(b.amt - self.tol_c, b.amt + self.tol_c + 1):
                for l in idx.get(a, []):
                    dd = self._dd(b, l)
                    if dd > w or self._ref_conflict(b, l):
                        continue
                    sim = similarity(b.ndesc, l.ndesc)
                    if sim >= self.cfg.desc_match_threshold:
                        c.append((0.7 * sim + 0.3 * (1 - dd / (w + 1)), b, l))
        self._assign(c, R2)

    def _r3_reference_mismatch(self):
        idx, c = self._ref_index(), []
        for b in self._free(self.bank):
            for l in idx.get(b.nref, []) if b.nref else []:
                dd = self._dd(b, l)
                if dd <= self.cfg.ref_date_window_days and not self._amt_eq(b, l):
                    c.append((0.95 - 0.005 * dd, b, l))
        self._assign(c, R3, "AMOUNT_MISMATCH", self._mismatch_subtype)

    def _r4_amount_date_only(self):
        idx, c = defaultdict(list), []
        for l in self._free(self.ledger):
            idx[l.amt].append(l)
        w = self.cfg.date_tolerance_days
        for b in self._free(self.bank):
            for a in range(b.amt - self.tol_c, b.amt + self.tol_c + 1):
                for l in idx.get(a, []):
                    dd = self._dd(b, l)
                    if dd <= w and not self._ref_conflict(b, l):
                        c.append((max(0.3, 0.6 - 0.05 * dd), b, l))
        self._assign(c, R4, "LOW_CONFIDENCE")

    def _r5_split(self):
        cfg, n = self.cfg, 0
        for single_is_bank in (True, False):
            pool = sorted(self._free(self.ledger if single_is_bank else self.bank),
                          key=lambda t: t.date)
            keys = [t.date.toordinal() for t in pool]
            for s in self._free(self.bank if single_is_bank else self.ledger):
                if s.uid in self.matched or s.amt == 0:
                    continue
                o, w = s.date.toordinal(), cfg.split_date_window_days
                lo, hi = bisect.bisect_left(keys, o - w), bisect.bisect_right(keys, o + w)
                cand = [p for p in pool[lo:hi]
                        if p.uid not in self.matched and p.amt * s.amt > 0
                        and abs(p.amt) < abs(s.amt)]
                if len(cand) < 2:
                    continue
                cand.sort(key=lambda p: (self._dd(p, s), p.uid))
                cand = cand[:cfg.split_candidate_cap]
                found = None
                for k in range(2, cfg.split_max_items + 1):
                    for combo in combinations(cand, k):
                        if abs(sum(p.amt for p in combo) - s.amt) <= self.tol_c:
                            found = combo
                            break
                    if found:
                        break
                if found:
                    b, l = ([s], list(found)) if single_is_bank else (list(found), [s])
                    self._commit(b, l, R5, 0.5, "SPLIT")
                    n += 1
        self.rule_log.append((R5, n))

    def _r6_fuzzy_mismatch(self):
        cfg, c = self.cfg, []
        by_date = defaultdict(list)
        for l in self._free(self.ledger):
            by_date[l.date].append(l)
        for b in self._free(self.bank):
            for off in range(-cfg.date_tolerance_days, cfg.date_tolerance_days + 1):
                for l in by_date.get(b.date + timedelta(days=off), []):
                    if l.uid in self.matched or self._amt_eq(b, l) or self._ref_conflict(b, l):
                        continue
                    sim = similarity(b.ndesc, l.ndesc)
                    if sim < cfg.mismatch_desc_threshold:
                        continue
                    if b.amt == -l.amt and b.amt != 0:
                        c.append((0.9 * sim, b, l))
                    elif b.amt * l.amt > 0 and (
                            abs(b.amt - l.amt) <= cfg.mismatch_max_pct * max(abs(b.amt), abs(l.amt))):
                        c.append((0.85 * sim, b, l))
        self._assign(c, R6, "AMOUNT_MISMATCH", self._mismatch_subtype)

    def _mismatch_subtype(self, b: Txn, l: Txn) -> str:
        diff = b.amt - l.amt
        if b.amt == -l.amt and b.amt != 0:
            return "SIGN_REVERSAL"
        if abs(diff) <= round(self.cfg.rounding_threshold * 100):
            return "ROUNDING_DIFFERENCE"
        if sorted(str(abs(b.amt))) == sorted(str(abs(l.amt))) and diff % 9 == 0:
            return "POSSIBLE_TRANSPOSITION"
        return "BANK_AMOUNT_LOWER" if abs(b.amt) < abs(l.amt) else "BANK_AMOUNT_HIGHER"

    # -- classification ----------------------------------------------------- #
    def _sev(self, amount_c: int, floor="Low", cap="High") -> str:
        a = abs(amount_c) / 100
        lvl = 2 if a >= self.cfg.high_materiality else 1 if a >= self.cfg.medium_materiality else 0
        lvl = min(max(lvl, LEVELS.index(floor)), LEVELS.index(cap))
        return LEVELS[lvl]

    def _mk(self, typ, sub, bank, ledger, sev, expl, action, match_id="") -> ExceptionItem:
        refs = [t.ref or t.nref for t in bank + ledger if (t.ref or t.nref)]
        if bank and ledger:
            desc = f"B: {bank[0].desc} | L: {ledger[0].desc}"
        else:
            desc = (bank or ledger)[0].desc
        return ExceptionItem(
            type=typ, subtype=sub, severity=sev,
            bank_ids="; ".join(t.uid for t in bank), ledger_ids="; ".join(t.uid for t in ledger),
            bank_date=max((t.date for t in bank), default=None),
            ledger_date=max((t.date for t in ledger), default=None),
            reference=refs[0] if refs else "", description=desc[:240],
            bank_amt=sum(t.amt for t in bank) if bank else None,
            ledger_amt=sum(t.amt for t in ledger) if ledger else None,
            explanation=expl, action=action, match_id=match_id)

    def _is_duplicate(self, t: Txn, idx: dict) -> Optional[Txn]:
        """Return an identical 'twin' in the same source, if t is the surplus copy."""
        for a in range(t.amt - self.tol_c, t.amt + self.tol_c + 1):
            for u in idx.get(a, []):
                if u is t or self._dd(t, u) > self.cfg.duplicate_window_days:
                    continue
                same = (t.nref and t.nref == u.nref) or similarity(t.ndesc, u.ndesc) >= 0.9
                if not same:
                    continue
                if u.uid in self.matched:
                    return u
                if (u.date, u.row) < (t.date, t.row):
                    return u
        return None

    def _classify(self):
        ex = self.exceptions
        # 1) match-based exceptions
        for m in self.matches:
            b, l = m.bank, m.ledger
            if m.kind == "AMOUNT_MISMATCH":
                ex.append(self._mismatch_exception(m))
            elif m.kind == "LOW_CONFIDENCE":
                ex.append(self._mk(
                    LOW_CONFIDENCE, "NO_REFERENCE_OR_DESCRIPTION_SUPPORT", b, l,
                    self._sev(b[0].amt, cap="Medium"),
                    f"Paired on amount {fm(b[0].amt)} and date (within "
                    f"{self.cfg.date_tolerance_days} days) only; no reference or description corroboration.",
                    "Confirm the pairing against source documents before accepting.", m.match_id))
            elif m.kind == "SPLIT":
                total = m.bank_total
                many = "ledger entries" if len(l) > 1 else "bank lines"
                ex.append(self._mk(
                    SPLIT_MATCH, "BATCH_DEPOSIT_OR_PAYMENT", b, l,
                    self._sev(total, cap="Medium"),
                    f"{max(len(b), len(l))} {many} sum exactly to {fm(abs(total))} on the other side "
                    f"(date window {self.cfg.split_date_window_days} days).",
                    "Confirm batch composition against deposit slip / remittance advice.", m.match_id))

        # 2) residual items: duplicates, then unmatched
        b_idx, l_idx = defaultdict(list), defaultdict(list)
        for t in self.bank:
            b_idx[t.amt].append(t)
        for t in self.ledger:
            l_idx[t.amt].append(t)

        for t in self._free(self.bank):
            twin = self._is_duplicate(t, b_idx)
            if twin:
                ex.append(self._mk(
                    DUPLICATE_BANK, "REPEATED_BANK_ENTRY", [t], [],
                    self._sev(t.amt, floor="Medium"),
                    f"Identical to {twin.uid} ({twin.date}, {fm(twin.amt)}); only one ledger entry exists.",
                    "Raise duplicate debit/credit claim with the bank and request reversal."))
                continue
            low = t.ndesc
            if re.search(r"\b(FEE|FEES|CHARGE|CHARGES|CHGS|INTEREST|INT|TAX|GST|COMMISSION)\b", low):
                sub, act = ("UNRECORDED_BANK_CHARGE_OR_INTEREST",
                            "Verify with bank advice, then post journal in the ledger.")
            elif t.amt >= 0:
                sub, act = ("UNRECORDED_RECEIPT",
                            "Identify the payer; record receipt / allocate to open invoice, or query the bank.")
            else:
                sub, act = ("UNRECORDED_PAYMENT",
                            "Obtain supporting document; record in ledger, or dispute if unauthorised.")
            ex.append(self._mk(
                UNMATCHED_BANK, sub, [t], [], self._sev(t.amt),
                f"Bank line of {fm(t.amt)} on {t.date} has no corresponding ledger entry.", act))

        for t in self._free(self.ledger):
            twin = self._is_duplicate(t, l_idx)
            if twin:
                ex.append(self._mk(
                    DUPLICATE_LEDGER, "DOUBLE_POSTED_ENTRY", [], [t],
                    self._sev(t.amt, floor="Medium"),
                    f"Identical to {twin.uid} ({twin.date}, {fm(twin.amt)}); only one bank line exists.",
                    "Verify against source document and reverse the duplicate ledger entry."))
                continue
            age = (self.period_end - t.date).days if self.period_end else 0
            if age <= self.cfg.outstanding_days:
                sub = "OUTSTANDING_IN_TRANSIT"
                act = "Timing difference - confirm it clears on the next statement; carry forward."
                sev = self._sev(t.amt, cap="Medium")
                note = f"Recorded {age} day(s) before statement end; likely still in transit."
            elif t.amt < 0:
                sub = "STALE_OUTSTANDING_PAYMENT"
                act = "Aged unpresented payment - confirm status with payee; consider stop-payment / reissue."
                sev = self._sev(t.amt, floor="Medium")
                note = f"Recorded {age} days before statement end and still not cleared."
            else:
                sub = "AGED_UNCLEARED_RECEIPT"
                act = "Receipt not credited by bank - verify deposit slip / remittance; escalate."
                sev = self._sev(t.amt, floor="Medium")
                note = f"Recorded {age} days before statement end and still not credited."
            ex.append(self._mk(
                UNMATCHED_LEDGER, sub, [], [t], sev,
                f"Ledger entry of {fm(t.amt)} on {t.date} has no bank line. {note}", act))

        # 3) order for the reviewer and assign IDs
        order = {t: i for i, t in enumerate(TYPE_ORDER)}
        ex.sort(key=lambda e: (-LEVELS.index(e.severity), order[e.type], -abs(e.impact), e.bank_ids, e.ledger_ids))
        for i, e in enumerate(ex, 1):
            e.exc_id = f"EX{i:04d}"
            for uid in (e.bank_ids + "; " + e.ledger_ids).split("; "):
                if uid.strip():
                    self.txn_exc[uid.strip()] = e.exc_id

    def _mismatch_exception(self, m: Match) -> ExceptionItem:
        b, l, sub = m.bank[0], m.ledger[0], m.subtype
        diff = b.amt - l.amt
        base = f"Bank {fm(b.amt)} vs ledger {fm(l.amt)} (difference {fm(diff)}). "
        text, floor, action = {
            "SIGN_REVERSAL": ("Amounts are equal and opposite; ledger entry was likely posted with the wrong sign.",
                              "Medium", "Correct the ledger entry sign (reverse and re-post)."),
            "ROUNDING_DIFFERENCE": ("Difference is within the rounding threshold; likely rounding or minor keying variance.",
                                    "Low", "Post a rounding adjustment if policy allows, else correct the entry."),
            "POSSIBLE_TRANSPOSITION": ("Digits are a permutation of each other and the difference is divisible by 9; "
                                       "likely a keying transposition error.",
                                       "Medium", "Check the source document; correct whichever side is wrong."),
            "BANK_AMOUNT_LOWER": ("Bank amount is smaller in magnitude than the ledger; possible bank fee deducted, "
                                  "short payment, or ledger over-recording.",
                                  "Low", "Obtain remittance / bank advice; post fee or raise short-payment query."),
            "BANK_AMOUNT_HIGHER": ("Bank amount is larger in magnitude than the ledger; possible unrecorded charge, "
                                   "over-payment, or ledger under-recording.",
                                   "Low", "Review source document; correct ledger or query the bank."),
        }[sub]
        return self._mk(AMOUNT_MISMATCH, sub, [b], [l], self._sev(diff, floor=floor),
                        base + text, action, m.match_id)

    # -- reporting helpers -------------------------------------------------- #
    def bridge(self) -> dict:
        """Reconciliation proof: net difference must equal the sum of explained items."""
        net = sum(t.amt for t in self.bank) - sum(t.amt for t in self.ledger)
        in_match = sum(m.bank_total - m.ledger_total for m in self.matches)
        by = defaultdict(int)
        for e in self.exceptions:
            if e.type in (UNMATCHED_BANK, UNMATCHED_LEDGER, DUPLICATE_BANK, DUPLICATE_LEDGER):
                by[e.type] += e.impact
        return {"net_difference": net, "matched_pair_differences": in_match, **by,
                "unexplained": net - in_match - sum(by.values())}


# --------------------------------------------------------------------------- #
# Excel export
# --------------------------------------------------------------------------- #
FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
SEC_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
SEV_FILL = {"High": PatternFill("solid", fgColor="F8CBAD"),
            "Medium": PatternFill("solid", fgColor="FFE699"),
            "Low": PatternFill("solid", fgColor="E2EFDA")}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
AMT, INT_, DATE_, PCT = "#,##0.00;(#,##0.00);-", "#,##0", "yyyy-mm-dd", "0.0%"


def _w(ws, r, c, v, fmt=None, bold=False, color="000000", fill=None, wrap=False,
       text=False, border=True, size=10, align=None):
    cell = ws.cell(r, c, v)
    if text and isinstance(v, str):
        cell.data_type = "s"          # never let imported text be evaluated as a formula
    cell.font = Font(name=FONT, size=size, bold=bold, color=color)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if border:
        cell.border = BORDER
    cell.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal=align)
    return cell


def _header(ws, row, headers):
    for j, h in enumerate(headers, 1):
        c = _w(ws, row, j, h, bold=True, color="FFFFFF", fill=HDR_FILL, align="center")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[row].height = 32


def _title(ws, title, note):
    _w(ws, 1, 1, title, bold=True, size=14, border=False)
    _w(ws, 2, 1, note, border=False, color="595959")


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def export_excel(rec: Reconciler, cfg: Config, meta: dict, path) -> None:
    wb = Workbook()
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_ex = wb.create_sheet("Exceptions")
    ws_m = wb.create_sheet("Matched")
    ws_b = wb.create_sheet("Bank_Txns")
    ws_l = wb.create_sheet("Ledger_Txns")
    ws_a = wb.create_sheet("Audit_Trail")

    nb, nl, ne, nm = len(rec.bank), len(rec.ledger), len(rec.exceptions), len(rec.matches)
    last = lambda n: max(5, 4 + n)

    # ---------------- Exceptions ------------------------------------------- #
    _title(ws_ex, "Exception Report - for auditor review and sign-off",
           "Yellow columns P-S are reviewer inputs. Difference = Bank - Ledger. "
           "Every exception needs a disposition before sign-off (see Summary).")
    ex_headers = ["Exception ID", "Type", "Sub-type", "Severity", "Bank Txn ID", "Ledger Txn ID",
                  "Bank Date", "Ledger Date", "Reference", "Description", "Bank Amount",
                  "Ledger Amount", "Difference (Bank - Ledger)", "Explanation (auto-generated)",
                  "Suggested Action", "Disposition", "Reviewer Comments", "Reviewed By", "Review Date"]
    _header(ws_ex, 4, ex_headers)
    for i, e in enumerate(rec.exceptions):
        r = 5 + i
        _w(ws_ex, r, 1, e.exc_id)
        _w(ws_ex, r, 2, e.type)
        _w(ws_ex, r, 3, e.subtype, wrap=True)
        _w(ws_ex, r, 4, e.severity, fill=SEV_FILL[e.severity], bold=True, align="center")
        _w(ws_ex, r, 5, e.bank_ids, wrap=True)
        _w(ws_ex, r, 6, e.ledger_ids, wrap=True)
        _w(ws_ex, r, 7, e.bank_date, DATE_)
        _w(ws_ex, r, 8, e.ledger_date, DATE_)
        _w(ws_ex, r, 9, e.reference, text=True)
        _w(ws_ex, r, 10, e.description, wrap=True, text=True)
        _w(ws_ex, r, 11, None if e.bank_amt is None else e.bank_amt / 100, AMT)
        _w(ws_ex, r, 12, None if e.ledger_amt is None else e.ledger_amt / 100, AMT)
        _w(ws_ex, r, 13, f"=N(K{r})-N(L{r})", AMT, bold=True)
        _w(ws_ex, r, 14, e.explanation, wrap=True)
        _w(ws_ex, r, 15, e.action, wrap=True)
        for c in range(16, 20):
            _w(ws_ex, r, c, None, fill=INPUT_FILL, color="0000FF", wrap=(c == 17),
               fmt=DATE_ if c == 19 else None)
    if ne:
        dv = DataValidation(type="list", formula1=f'"{DISPOSITIONS}"', allow_blank=True)
        dv.error, dv.errorTitle = "Pick a disposition from the list.", "Invalid disposition"
        ws_ex.add_data_validation(dv)
        dv.add(f"P5:P{4 + ne}")
        ws_ex.auto_filter.ref = f"A4:S{4 + ne}"
    ws_ex.freeze_panes = "C5"
    _widths(ws_ex, [11, 20, 24, 10, 12, 14, 12, 12, 12, 42, 14, 14, 16, 58, 44, 22, 34, 16, 13])

    # ---------------- Matched ---------------------------------------------- #
    _title(ws_m, "Matched Transactions (audit log of every pairing)",
           "Dates shown are the latest date in a group. Flag column links to the exception raised, if any.")
    _header(ws_m, 4, ["Match ID", "Rule", "Confidence", "Bank Txn ID(s)", "Ledger Txn ID(s)",
                      "Bank Date", "Ledger Date", "Date Variance (days)", "Bank Amount",
                      "Ledger Amount", "Difference", "Bank Description", "Ledger Description",
                      "Exception Flag"])
    exc_by_match = {e.match_id: e.exc_id for e in rec.exceptions if e.match_id}
    for i, m in enumerate(rec.matches):
        r = 5 + i
        _w(ws_m, r, 1, m.match_id)
        _w(ws_m, r, 2, m.rule)
        _w(ws_m, r, 3, m.score, "0.00")
        _w(ws_m, r, 4, "; ".join(t.uid for t in m.bank), wrap=True)
        _w(ws_m, r, 5, "; ".join(t.uid for t in m.ledger), wrap=True)
        _w(ws_m, r, 6, max(t.date for t in m.bank), DATE_)
        _w(ws_m, r, 7, max(t.date for t in m.ledger), DATE_)
        _w(ws_m, r, 8, f"=ABS(F{r}-G{r})", "0")
        _w(ws_m, r, 9, m.bank_total / 100, AMT)
        _w(ws_m, r, 10, m.ledger_total / 100, AMT)
        _w(ws_m, r, 11, f"=I{r}-J{r}", AMT)
        _w(ws_m, r, 12, m.bank[0].desc, text=True)
        _w(ws_m, r, 13, m.ledger[0].desc, text=True)
        _w(ws_m, r, 14, exc_by_match.get(m.match_id, ""))
    if nm:
        ws_m.auto_filter.ref = f"A4:N{4 + nm}"
    ws_m.freeze_panes = "B5"
    _widths(ws_m, [10, 34, 11, 16, 18, 12, 12, 12, 14, 14, 13, 38, 38, 12])

    # ---------------- Bank / Ledger transaction sheets --------------------- #
    for ws, txns, name in ((ws_b, rec.bank, "Bank"), (ws_l, rec.ledger, "Ledger")):
        _title(ws, f"{name} Transactions - full population with reconciliation status",
               "Status: MATCHED = clean, MATCHED-REVIEW = paired but flagged, UNMATCHED = exception raised.")
        _header(ws, 4, ["Txn ID", "Source Row", "Date", "Description", "Reference", "Amount",
                        "Status", "Match ID", "Exception ID"])
        for i, t in enumerate(txns):
            r, m = 5 + i, rec.matched.get(t.uid)
            st = rec.status(t)
            _w(ws, r, 1, t.uid)
            _w(ws, r, 2, t.row, "0")
            _w(ws, r, 3, t.date, DATE_)
            _w(ws, r, 4, t.desc, text=True)
            _w(ws, r, 5, t.ref, text=True)
            _w(ws, r, 6, t.amt / 100, AMT)
            _w(ws, r, 7, st, fill=SEV_FILL["Low"] if st == "MATCHED" else
               SEV_FILL["Medium"] if st == "MATCHED-REVIEW" else SEV_FILL["High"])
            _w(ws, r, 8, m.match_id if m else "")
            _w(ws, r, 9, rec.txn_exc.get(t.uid, ""))
        if txns:
            ws.auto_filter.ref = f"A4:I{4 + len(txns)}"
        ws.freeze_panes = "B5"
        _widths(ws, [10, 11, 12, 46, 16, 15, 17, 10, 12])

    # ---------------- Summary ---------------------------------------------- #
    BK = lambda c: f"Bank_Txns!${c}$5:${c}${last(nb)}"
    LD = lambda c: f"Ledger_Txns!${c}$5:${c}${last(nl)}"
    EX = lambda c: f"Exceptions!${c}$5:${c}${last(ne)}"
    MT = lambda c: f"Matched!${c}$5:${c}${last(nm)}"
    ws = ws_sum
    _w(ws, 1, 1, "Bank Reconciliation Summary", bold=True, size=16, border=False)
    _w(ws, 2, 1, f"Run {meta['run_id']}  |  generated {meta['started_at']} UTC  |  "
                 f"bank_reconciler v{__version__}", border=False, color="595959")
    row = [4]

    def section(title):
        for c in range(1, 7):
            _w(ws, row[0], c, title if c == 1 else None, bold=True, fill=SEC_FILL)
        row[0] += 1

    def line(label, value, fmt=None, bold=False):
        _w(ws, row[0], 1, label, bold=bold)
        _w(ws, row[0], 2, value, fmt, bold=bold)
        ref = f"B{row[0]}"
        row[0] += 1
        return ref

    section("1. Population")
    bank_n = line("Bank transactions loaded", f"=COUNTA({BK('A')})", INT_)
    led_n = line("Ledger transactions loaded", f"=COUNTA({LD('A')})", INT_)
    line("Bank statement - first date", f"=MIN({BK('C')})", DATE_)
    line("Bank statement - last date", f"=MAX({BK('C')})", DATE_)
    bank_net = line("Bank net movement", f"=SUM({BK('F')})", AMT)
    led_net = line("Ledger net movement", f"=SUM({LD('F')})", AMT)
    net_diff = line("Net difference (Bank - Ledger)", f"={bank_net}-{led_net}", AMT, bold=True)
    row[0] += 1

    section("2. Matching results")
    b_ok = line("Bank - clean matches", f'=COUNTIF({BK("G")},"MATCHED")', INT_)
    b_rv = line("Bank - matched but flagged for review", f'=COUNTIF({BK("G")},"MATCHED-REVIEW")', INT_)
    b_un = line("Bank - unmatched / duplicate", f'=COUNTIF({BK("G")},"UNMATCHED")', INT_)
    line("Bank auto-match rate", f"=IF({bank_n}=0,0,({b_ok}+{b_rv})/{bank_n})", PCT, True)
    l_ok = line("Ledger - clean matches", f'=COUNTIF({LD("G")},"MATCHED")', INT_)
    l_rv = line("Ledger - matched but flagged for review", f'=COUNTIF({LD("G")},"MATCHED-REVIEW")', INT_)
    l_un = line("Ledger - unmatched / duplicate", f'=COUNTIF({LD("G")},"UNMATCHED")', INT_)
    line("Ledger auto-match rate", f"=IF({led_n}=0,0,({l_ok}+{l_rv})/{led_n})", PCT, True)
    line("Match groups created", f"=COUNTA({MT('A')})", INT_)
    line("Matches with date variance > 0 days (informational)", f'=COUNTIF({MT("H")},">0")', INT_)
    row[0] += 1

    section("3. Exceptions by type")
    for j, h in enumerate(["Exception type", "Count", "Net impact (Bank - Ledger)",
                           "High", "Medium", "Low"], 1):
        c = _w(ws, row[0], j, h, bold=True, color="FFFFFF", fill=HDR_FILL)
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    row[0] += 1
    t_first = row[0]
    for typ in TYPE_ORDER:
        r = row[0]
        _w(ws, r, 1, typ)
        _w(ws, r, 2, f"=COUNTIF({EX('B')},$A{r})", INT_)
        _w(ws, r, 3, f"=SUMIFS({EX('M')},{EX('B')},$A{r})", AMT)
        for j, sev in enumerate(("High", "Medium", "Low"), 4):
            _w(ws, r, j, f'=COUNTIFS({EX("B")},$A{r},{EX("D")},"{sev}")', INT_)
        row[0] += 1
    t_last = row[0] - 1
    r = row[0]
    _w(ws, r, 1, "Total", bold=True)
    for j in range(2, 7):
        col = get_column_letter(j)
        _w(ws, r, j, f"=SUM({col}{t_first}:{col}{t_last})", AMT if j == 3 else INT_, bold=True)
    exc_total = f"B{r}"
    row[0] += 2

    section("4. Reconciliation bridge (control proof)")
    nd = line("Net difference (Bank - Ledger)", f"={net_diff}", AMT, True)
    first = row[0]
    line("Less: amount differences inside matched pairs", f"=SUM({MT('K')})", AMT)
    for label, typ in (("Less: unmatched bank items", UNMATCHED_BANK),
                       ("Less: duplicate bank items", DUPLICATE_BANK),
                       ("Add back: unmatched ledger items", UNMATCHED_LEDGER),
                       ("Add back: duplicate ledger items", DUPLICATE_LEDGER)):
        line(label, f'=SUMIFS({EX("M")},{EX("B")},"{typ}")', AMT)
    lastb = row[0] - 1
    unexpl = line("Unexplained difference (must be 0.00)", f"={nd}-SUM(B{first}:B{lastb})", AMT, True)
    chk = line("Control check", f'=IF(ABS({unexpl})<0.005,"PASS - fully explained","FAIL - unexplained difference")', None, True)
    ws.conditional_formatting.add(chk, FormulaRule(formula=[f'LEFT({chk},4)="PASS"'],
                                  fill=SEV_FILL["Low"]))
    ws.conditional_formatting.add(chk, FormulaRule(formula=[f'LEFT({chk},4)="FAIL"'],
                                  fill=SEV_FILL["High"]))
    row[0] += 1

    section("5. Review status")
    disp = line("Exceptions with a disposition", f'=COUNTIF({EX("P")},"?*")', INT_)
    opn = line("Exceptions still open", f"={exc_total}-{disp}", INT_, True)
    line("High-severity exceptions still open",
         f'=COUNTIF({EX("D")},"High")-COUNTIFS({EX("D")},"High",{EX("P")},"?*")', INT_)
    status = line("Sign-off status",
                  f'=IF(AND(ABS({unexpl})<0.005,{opn}=0),"READY FOR SIGN-OFF",'
                  f'"NOT READY - "&{opn}&" open exception(s)"&IF(ABS({unexpl})>=0.005," + unexplained difference",""))',
                  None, True)
    ws.conditional_formatting.add(status, FormulaRule(formula=[f'LEFT({status},5)="READY"'],
                                  fill=SEV_FILL["Low"]))
    ws.conditional_formatting.add(status, FormulaRule(formula=[f'LEFT({status},3)="NOT"'],
                                  fill=SEV_FILL["High"]))
    row[0] += 1

    section("6. Sign-off")
    for j, h in enumerate(["Role", "Name", "Signature", "Date"], 1):
        _w(ws, row[0], j, h, bold=True, color="FFFFFF", fill=HDR_FILL, align="center")
    row[0] += 1
    for role, name, d in (("Prepared by (system)", meta["user"], meta["started_at"][:10]),
                          ("Reviewed by", None, None), ("Approved by", None, None)):
        _w(ws, row[0], 1, role)
        for j, v in enumerate((name, None, d), 2):
            _w(ws, row[0], j, v, fill=INPUT_FILL, color="0000FF", text=True)
        row[0] += 1
    _widths(ws, [54, 22, 26, 12, 12, 12])
    ws.sheet_view.showGridLines = False

    # ---------------- Audit trail ------------------------------------------ #
    wa = ws_a
    _w(wa, 1, 1, "Audit Trail - provenance of this reconciliation", bold=True, size=14, border=False)
    _w(wa, 2, 1, "Input hashes, parameters and rule counts. The hash of this workbook is recorded in the "
                 "hash-chained audit_log.jsonl written beside it.", border=False, color="595959")
    r = [4]

    def a_section(t):
        for c in range(1, 8):
            _w(wa, r[0], c, t if c == 1 else None, bold=True, fill=SEC_FILL)
        r[0] += 1

    def kv(k, v):
        _w(wa, r[0], 1, k, bold=True)
        _w(wa, r[0], 2, v, text=True, wrap=True)
        r[0] += 1

    a_section("Run")
    for k in ("run_id", "started_at", "user", "host", "tool_version", "python", "pandas"):
        kv(k, str(meta[k]))
    r[0] += 1
    a_section("Input files")
    for j, h in enumerate(["Source", "File", "Bytes", "SHA-256", "Rows read", "Rows loaded", "Rows rejected"], 1):
        _w(wa, r[0], j, h, bold=True, color="FFFFFF", fill=HDR_FILL)
    r[0] += 1
    for f in meta["inputs"]:
        for j, k in enumerate(["source", "file", "bytes", "sha256", "rows_read", "rows_loaded", "rows_rejected"], 1):
            _w(wa, r[0], j, f[k], text=True)
        r[0] += 1
    r[0] += 1
    a_section("Matching parameters")
    for k, v in asdict(cfg).items():
        kv(k, v)
    r[0] += 1
    a_section("Rule execution log (in order)")
    for rule, n in rec.rule_log:
        kv(rule, n)
    r[0] += 1
    a_section("Data-quality log: rows rejected on load (never silently dropped)")
    if not meta["rejects"]:
        kv("Rejected rows", "None")
    for j, h in enumerate(["Source", "Source row", "Reason", "Raw values"], 1):
        if meta["rejects"]:
            _w(wa, r[0], j, h, bold=True, color="FFFFFF", fill=HDR_FILL)
    if meta["rejects"]:
        r[0] += 1
    for rej in meta["rejects"][:500]:
        for j, k in enumerate(["source", "row", "reason", "raw"], 1):
            _w(wa, r[0], j, rej[k], text=True)
        r[0] += 1
    _widths(wa, [34, 50, 14, 66, 12, 12, 14])

    wb.save(path)


# --------------------------------------------------------------------------- #
# Tamper-evident audit log (hash chain)
# --------------------------------------------------------------------------- #
def _entry_hash(entry: dict) -> str:
    body = {k: v for k, v in entry.items() if k != "entry_hash"}
    return hashlib.sha256(json.dumps(body, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def append_audit_log(path: Path, entry: dict) -> None:
    prev = "GENESIS"
    if path.exists():
        last = None
        with open(path, encoding="utf-8") as f:
            for ln in f:
                if ln.strip():
                    last = ln
        if last:
            prev = json.loads(last)["entry_hash"]
    entry["prev_hash"] = prev
    entry["entry_hash"] = _entry_hash(entry)
    with open(path, "a", encoding="utf-8") as f:
        f.write(json.dumps(entry, ensure_ascii=False) + "\n")


def verify_audit_log(path: Path) -> tuple[bool, str]:
    prev, n = "GENESIS", 0
    with open(path, encoding="utf-8") as f:
        for i, ln in enumerate(f, 1):
            if not ln.strip():
                continue
            e = json.loads(ln)
            if e.get("prev_hash") != prev:
                return False, f"Chain broken at entry {i}: prev_hash mismatch (entry removed or reordered)."
            if e.get("entry_hash") != _entry_hash(e):
                return False, f"Entry {i} was modified: hash mismatch."
            prev, n = e["entry_hash"], n + 1
    return True, f"Audit log intact: {n} entr{'y' if n == 1 else 'ies'} verified."


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def _kv_map(s: Optional[str]) -> dict:
    out = {}
    for part in (s or "").split(","):
        if "=" in part:
            k, v = part.split("=", 1)
            out[k.strip().lower()] = v.strip()
    return out


def _sheet(v):
    return int(v) if str(v).isdigit() else v


def build_parser():
    p = argparse.ArgumentParser(
        description="Automated bank reconciliation: match, classify, export audit-ready Excel.")
    p.add_argument("--bank", help="Bank statement (.csv/.xlsx)")
    p.add_argument("--ledger", help="Accounting records / cash book (.csv/.xlsx)")
    p.add_argument("--out", default="reconciliation_report.xlsx")
    p.add_argument("--bank-cols", help='Column overrides, e.g. "date=Posting Date,amount=Amt"')
    p.add_argument("--ledger-cols", help="Same as --bank-cols for the ledger file")
    p.add_argument("--bank-sheet", default="0")
    p.add_argument("--ledger-sheet", default="0")
    p.add_argument("--config", help="JSON file overriding Config fields")
    p.add_argument("--audit-log", help="Path for hash-chained audit log (default: beside --out)")
    p.add_argument("--no-audit-log", action="store_true")
    p.add_argument("--verify-log", metavar="PATH", help="Verify an audit log's hash chain and exit")
    for fld, val in asdict(Config()).items():
        flag = "--" + fld.replace("_", "-")
        if isinstance(val, bool):
            p.add_argument(flag, action="store_true", default=None)
        else:
            p.add_argument(flag, type=type(val), default=None)
    return p


def main(argv=None) -> int:
    args = build_parser().parse_args(argv)
    if args.verify_log:
        ok, msg = verify_audit_log(Path(args.verify_log))
        print(msg)
        return 0 if ok else 1
    if not (args.bank and args.ledger):
        print("error: --bank and --ledger are required", file=sys.stderr)
        return 2

    cfg = Config()
    overrides = json.load(open(args.config)) if args.config else {}
    for k, v in overrides.items():
        if not hasattr(cfg, k):
            print(f"error: unknown config key '{k}'", file=sys.stderr)
            return 2
        setattr(cfg, k, v)
    for fld in asdict(cfg):
        v = getattr(args, fld, None)
        if v is not None:
            setattr(cfg, fld, v)

    try:
        bank, brej, binfo = load_transactions(args.bank, "BANK", cfg, _kv_map(args.bank_cols),
                                              _sheet(args.bank_sheet))
        ledger, lrej, linfo = load_transactions(args.ledger, "LEDGER", cfg, _kv_map(args.ledger_cols),
                                                _sheet(args.ledger_sheet))
    except (ValueError, FileNotFoundError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2

    rec = Reconciler(bank, ledger, cfg).run()
    meta = {"run_id": uuid.uuid4().hex[:12].upper(),
            "started_at": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S"),
            "user": getpass.getuser(), "host": platform.node(), "tool_version": __version__,
            "python": platform.python_version(), "pandas": pd.__version__,
            "inputs": [binfo, linfo], "rejects": brej + lrej}
    out = Path(args.out)
    export_excel(rec, cfg, meta, out)

    if not args.no_audit_log:
        log_path = Path(args.audit_log) if args.audit_log else out.parent / "audit_log.jsonl"
        append_audit_log(log_path, {
            "run_id": meta["run_id"], "timestamp_utc": meta["started_at"], "user": meta["user"],
            "tool_version": __version__, "inputs": meta["inputs"], "config": asdict(cfg),
            "results": {"bank_txns": len(bank), "ledger_txns": len(ledger),
                        "match_groups": len(rec.matches), "exceptions": rec.counts_by_type(),
                        "unexplained_cents": rec.bridge()["unexplained"],
                        "rows_rejected": len(meta["rejects"])},
            "output_file": out.name, "output_sha256": sha256_file(out)})

    br = rec.bridge()
    print(f"Run {meta['run_id']}  |  bank {len(bank)} txns  |  ledger {len(ledger)} txns")
    print(f"Match groups: {len(rec.matches)}  |  Exceptions: {len(rec.exceptions)}")
    for t in TYPE_ORDER:
        n = rec.counts_by_type().get(t, 0)
        if n:
            print(f"  {t:<24}{n:>4}")
    print(f"Unexplained difference: {fm(br['unexplained'])}   "
          f"{'(PASS)' if br['unexplained'] == 0 else '(FAIL)'}")
    if meta["rejects"]:
        print(f"WARNING: {len(meta['rejects'])} row(s) rejected on load - see Audit_Trail sheet.")
    print(f"Report written: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
